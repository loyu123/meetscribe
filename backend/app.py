"""语音转写 + 说话人分离 · 服务端（方案 B：异步任务化）。

形态：客户端(APP/网页) -> REST /api/v1 -> 分片上传 + 任务队列 -> 后台 worker 分段处理。

主要接口：
  POST /api/v1/uploads/init                 初始化分片上传
  POST /api/v1/uploads/<uid>/part           上传单个分片
  POST /api/v1/uploads/<uid>/complete       完成上传 -> 创建录音与任务
  POST /api/v1/uploads/simple               小文件一次性上传
  GET  /api/v1/jobs/<job_id>                任务状态 / 进度 / 分片明细
  POST /api/v1/jobs/<job_id>/retry          重跑失败分片
  GET  /api/v1/recordings                   录音列表
  GET  /api/v1/recordings/<id>              录音详情（含说话人分离结果）
  GET  /api/v1/recordings/<id>/media        播放音频（支持 Range 拖动）
  DELETE /api/v1/recordings/<id>            删除（含结果与音频）
  POST /api/v1/transcribe                   同步转写（仅适合短音频，快速自测用）
"""
import glob
import hashlib
import hmac
import json
import os
import re
import shutil
import tempfile
import threading
import time
import uuid
from urllib.parse import quote

from flask import Flask, jsonify, request, send_file, Response, g

import config
import polish
import rewrite
import store
import worker
import summary
import search
import ask
import gateway_client as gw
import ratelimit
import runtime_settings

config.ensure_dirs()
store.init_db()
store.ai_job_cleanup_stale()  # 重启后清理残留的 processing 纪要/润色任务（P0-1）
search.init_search_db()

app = Flask(__name__)
_START_TIME = time.time()  # P2-2：进程启动时刻，供 /system/metrics 上报运行时长
PART_SIZE = 4 * 1024 * 1024


def _json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


# ---------------- 鉴权与租户隔离（FR-31 / FR-08，opt-in） ----------------
def _current_tenant():
    return getattr(g, "tenant", None)


def _engine():
    """懒加载 engine（D2）：engine 依赖 torch/funasr，测试与 CI 导入 app 时不应拉起推理栈。"""
    import engine
    return engine


def _tenant_can_access(rid):
    """当前请求是否可访问该录音：未配置鉴权或管理员(*) 放行；
    否则仅自身租户或共享(tenant_id='')录音可见。"""
    t = _current_tenant()
    if t is None or t == "*":
        return True
    rec = store.get_recording(rid)
    if not rec:
        return False
    rt = rec.get("tenant_id") or ""
    return rt == t or rt == ""


def _require_admin():
    """P2-1：全局管理类操作（网关设置/全量重建/检索日志等）仅管理员或未鉴权模式可用。"""
    t = _current_tenant()
    if t and t != "*":
        return jsonify({"error": "仅管理员可执行此操作"}), 403
    return None


def _media_sign(rid):
    exp = int(time.time()) + config.MEDIA_SIGN_TTL
    msg = ("%s|%d" % (rid, exp)).encode("utf-8")
    sig = hmac.new(config.MEDIA_SIGN_SECRET.encode("utf-8"), msg, hashlib.sha256).hexdigest()
    return "%s.%d" % (sig, exp)


def _media_allowed(rid):
    """媒体下载是否放行：未配置签名密钥则放行（向后兼容）；
    配置后需带有效且未过期的 ?token= 签名。"""
    if not config.MEDIA_SIGN_SECRET:
        return True
    tok = request.args.get("token", "")
    try:
        sig, exp = tok.split(".")
        exp = int(exp)
    except Exception:
        return False
    if int(time.time()) > exp:
        return False
    msg = ("%s|%d" % (rid, exp)).encode("utf-8")
    exp_sig = hmac.new(config.MEDIA_SIGN_SECRET.encode("utf-8"), msg, hashlib.sha256).hexdigest()
    return hmac.compare_digest(sig, exp_sig)


def _media_url(rid):
    """带签名（若启用）的媒体地址，供前端直接播放/下载（FR-05）。"""
    if config.MEDIA_SIGN_SECRET:
        return "/api/v1/recordings/%s/media?token=%s" % (rid, _media_sign(rid))
    return "/api/v1/recordings/%s/media" % rid


@app.before_request
def _enforce_auth():
    # 未配置 ASR_API_KEYS 时不鉴权（向后兼容，不破坏现有前端/服务）
    if not config.API_KEY_TENANTS:
        return
    p = request.path
    if p == "/" or p.startswith("/static"):
        return
    auth = request.headers.get("Authorization", "")
    token = auth[len("Bearer "):] if auth.startswith("Bearer ") else auth
    tenant = config.API_KEY_TENANTS.get(token.strip())
    if tenant is None:
        return jsonify({"error": "unauthorized"}), 401
    g.tenant = tenant


# ---------------- 根路径 ----------------
@app.get("/")
def index():
    # 旧版内置 HTML 界面已移除，统一使用独立前端应用
    return jsonify({
        "service": "ASR 语音转写服务",
        "api": "/api/v1",
        "ui": "请使用前端应用（开发环境默认 http://localhost:5173）",
    })


# ---------------- 分片上传 ----------------
@app.post("/api/v1/uploads/init")
def upload_init():
    data = request.get_json(force=True, silent=True) or {}
    filename = (data.get("filename") or "audio").strip()
    ext = os.path.splitext(filename)[1].lower()
    if ext and ext not in config.ALLOWED_EXT:
        return jsonify({"error": "不支持的格式: %s" % ext}), 400
    uid = uuid.uuid4().hex
    d = os.path.join(config.UPLOAD_DIR, uid)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "meta.json"), "w", encoding="utf-8") as f:
        json.dump({"filename": filename, "ext": ext}, f)
    return jsonify({"upload_id": uid, "part_size": PART_SIZE})


@app.post("/api/v1/uploads/<uid>/part")
def upload_part(uid):
    if not re.fullmatch(r"[0-9a-f]{32}", uid):
        return jsonify({"error": "非法 upload_id"}), 400
    d = os.path.join(config.UPLOAD_DIR, uid)
    if not os.path.isdir(d):
        return jsonify({"error": "upload 不存在"}), 404
    f = request.files.get("file")
    if f is None:
        return jsonify({"error": "缺少分片"}), 400
    try:
        idx = int(request.form.get("index", "0"))
    except (TypeError, ValueError):
        return jsonify({"error": "index 非法"}), 400
    f.save(os.path.join(d, "part_%05d" % idx))
    return jsonify({"ok": True, "index": idx})


@app.post("/api/v1/uploads/<uid>/complete")
def upload_complete(uid):
    if not re.fullmatch(r"[0-9a-f]{32}", uid):
        return jsonify({"error": "非法 upload_id"}), 400
    d = os.path.join(config.UPLOAD_DIR, uid)
    if not os.path.isdir(d):
        return jsonify({"error": "upload 不存在"}), 404
    meta = _json(os.path.join(d, "meta.json")) or {}
    filename = meta.get("filename", "audio")
    ext = meta.get("ext", "")
    parts = sorted(glob.glob(os.path.join(d, "part_*")))
    if not parts:
        return jsonify({"error": "没有已上传的分片"}), 400

    # 数据完整性：分片序号必须 0..N-1 连续无缺口，否则拼接会产生静默损坏的音频
    idxs = [int(os.path.basename(p).split("_")[1]) for p in parts]
    if idxs != list(range(len(idxs))):
        shutil.rmtree(d, ignore_errors=True)
        return jsonify({"error": "分片不完整（存在缺失），请重新上传", "code": "INCOMPLETE_PARTS"}), 400

    os.makedirs(config.AUDIO_DIR, exist_ok=True)
    rid = uuid.uuid4().hex
    jid = uuid.uuid4().hex
    final = os.path.join(config.AUDIO_DIR, rid + ext)
    with open(final, "wb") as out:
        for p in parts:
            with open(p, "rb") as f:
                shutil.copyfileobj(f, out, 1024 * 1024)
    nbytes = os.path.getsize(final)
    # 服务端总大小上限（FR 前端 4GB 之外的服务端兜底）
    max_bytes = getattr(config, "MAX_UPLOAD_BYTES", 0)
    if max_bytes and nbytes > max_bytes:
        try:
            os.remove(final)
        except OSError:
            pass
        shutil.rmtree(d, ignore_errors=True)
        return jsonify({"error": "文件超过大小上限（%d GB）" % (max_bytes // 1024 ** 3), "code": "TOO_LARGE"}), 413

    data = request.get_json(force=True, silent=True) or {}
    hotword = (data.get("hotword") or "").strip()
    preset = data.get("preset_spk_num") or data.get("num_speakers")
    preset = int(preset) if preset and int(preset) > 0 else None
    callback = data.get("callback_url")
    asr_model = (data.get("asr_model") or config.DEFAULT_ASR_BACKEND)
    denoise = 1 if str(data.get("denoise", "0")) in ("1", "true", "True") else 0

    # FR-03 上传校验：解码元信息，损坏/无法解码/时长为0 直接拒绝，不进识别队列
    info = _engine().probe_audio(final)
    if not info["ok"]:
        try:
            os.remove(final)
        except OSError:
            pass
        shutil.rmtree(d, ignore_errors=True)
        return jsonify({"error": "音频文件损坏或无法解码（缺少音轨或时长为 0）", "code": "INVALID_AUDIO"}), 400

    # FR-06 幂等键：相同键重复提交直接返回已有任务，不重复执行/不计费
    idem = (data.get("idempotency_key") or "").strip()
    if idem:
        with _IDEM_LOCK:
            ex = store.get_idempotency(idem)
            # 幂等命中需校验录音仍存在；若已被删除则为陈旧映射，丢弃后按新上传处理
            if ex and store.get_recording(ex["recording_id"]):
                try:
                    os.remove(final)
                except OSError:
                    pass
                shutil.rmtree(d, ignore_errors=True)
                return jsonify({"recording_id": ex["recording_id"], "job_id": ex["job_id"], "idempotent": True})
            if ex:
                store.delete_idempotency(idem)
            store.save_idempotency(idem, rid, jid)  # 先占位，防止并发 complete 双建

    t = _current_tenant()
    tid = t if (t and t != "*") else ""
    # FR-08 租户配额：上传即校验当日处理时长额度（opt-in，未配 API_KEYS 不生效）
    ok, why = ratelimit.can_enqueue(tid, (info.get("duration_ms") or 0) / 1000.0)
    if not ok:
        try:
            os.remove(final)
        except OSError:
            pass
        shutil.rmtree(d, ignore_errors=True)
        return jsonify({"error": why, "code": "QUOTA_EXCEEDED"}), 429
    store.create_recording(rid, filename, ext, final, nbytes, num_speakers=preset, tenant_id=tid)
    store.create_job(jid, rid, hotword, preset, callback, asr_model, denoise)
    shutil.rmtree(d, ignore_errors=True)
    return jsonify({"recording_id": rid, "job_id": jid})


@app.post("/api/v1/uploads/simple")
def upload_simple():
    f = request.files.get("audio")
    if f is None or f.filename == "":
        return jsonify({"error": "未收到音频文件"}), 400
    filename = f.filename
    ext = os.path.splitext(filename)[1].lower() or ".wav"
    if ext not in config.ALLOWED_EXT:
        return jsonify({"error": "不支持的格式: %s" % ext}), 400
    os.makedirs(config.AUDIO_DIR, exist_ok=True)
    rid = uuid.uuid4().hex
    orig = os.path.join(config.AUDIO_DIR, rid + ext)
    f.save(orig)
    nbytes = os.path.getsize(orig)

    hotword = (request.form.get("hotword") or "").strip()
    preset = request.form.get("preset_spk_num") or request.form.get("num_speakers")
    preset = int(preset) if preset and int(preset) > 0 else None
    callback = request.form.get("callback_url")
    asr_model = request.form.get("asr_model") or config.DEFAULT_ASR_BACKEND
    denoise = 1 if str(request.form.get("denoise", "0")) in ("1", "true", "True") else 0

    # FR-03 上传校验
    info = _engine().probe_audio(orig)
    if not info["ok"]:
        try:
            os.remove(orig)
        except OSError:
            pass
        return jsonify({"error": "音频文件损坏或无法解码（缺少音轨或时长为 0）", "code": "INVALID_AUDIO"}), 400

    # FR-06 幂等键（与分片 complete 相同的加锁语义，防并发双建）
    jid = uuid.uuid4().hex
    idem = (request.form.get("idempotency_key") or "").strip()
    if idem:
        with _IDEM_LOCK:
            ex = store.get_idempotency(idem)
            # 幂等命中需校验录音仍存在；若已被删除则为陈旧映射，丢弃后按新上传处理
            if ex and store.get_recording(ex["recording_id"]):
                try:
                    os.remove(orig)
                except OSError:
                    pass
                return jsonify({"recording_id": ex["recording_id"], "job_id": ex["job_id"], "idempotent": True})
            if ex:
                store.delete_idempotency(idem)
            store.save_idempotency(idem, rid, jid)  # 先占位，防并发双建

    t = _current_tenant()
    tid = t if (t and t != "*") else ""
    # FR-08 租户配额：上传即校验当日处理时长额度（opt-in，未配 API_KEYS 不生效）
    dur_s = 0
    try:
        dur_s = float((_engine().probe_duration_ms(orig) or 0)) / 1000.0
    except Exception:
        dur_s = 0
    ok, why = ratelimit.can_enqueue(tid, dur_s)
    if not ok:
        try:
            os.remove(orig)
        except OSError:
            pass
        return jsonify({"error": why, "code": "QUOTA_EXCEEDED"}), 429
    store.create_recording(rid, filename, ext, orig, nbytes, num_speakers=preset, tenant_id=tid)
    store.create_job(jid, rid, hotword, preset, callback, asr_model, denoise)
    return jsonify({"recording_id": rid, "job_id": jid})


# ---------------- 任务 ----------------
@app.get("/api/v1/jobs/<jid>")
def job_status(jid):
    job = store.get_job(jid)
    if not job:
        return jsonify({"error": "任务不存在"}), 404
    if not _tenant_can_access(job.get("recording_id")):
        return jsonify({"error": "任务不存在"}), 404
    chunks = [{
        "idx": c["idx"], "start_ms": c["start_ms"], "end_ms": c["end_ms"],
        "status": c["status"], "attempts": c["attempts"], "error": c["error"],
    } for c in store.get_chunks(jid)]
    return jsonify({"job": job, "chunks": chunks})


@app.post("/api/v1/jobs/<jid>/retry")
def job_retry(jid):
    job = store.get_job(jid)
    if not job:
        return jsonify({"error": "任务不存在"}), 404
    if not _tenant_can_access(job.get("recording_id")):
        return jsonify({"error": "任务不存在"}), 404
    # 仅失败/部分完成的任务允许重试；processing 中重置会导致两个 worker 并发处理同一任务
    if job.get("status") not in ("failed", "partial"):
        return jsonify({"error": "当前状态（%s）不可重试" % job.get("status")}), 400
    requeued = worker.retry_job(jid)
    return jsonify({"ok": True, "requeued": requeued})


# ---------------- 纪要模板（内置 + 自定义，P2-1 租户隔离） ----------------
def _tenant_can_manage_template(tpl):
    """管理员/未鉴权可管理全部；普通租户仅可管理自己创建的（共享 '' 模板只读）。"""
    t = _current_tenant()
    if not t or t == "*":
        return True
    return tpl.get("tenant_id") == t


@app.get("/api/v1/templates")
def templates_list():
    t = _current_tenant()
    tenant_param = t if (t and t != "*") else None
    return jsonify({"templates": store.list_templates(tenant_id=tenant_param)})


@app.post("/api/v1/templates")
def templates_create():
    data = request.get_json(force=True, silent=True) or {}
    t = _current_tenant()
    tid, err = store.create_template(
        data.get("name"), data.get("color"), data.get("description"),
        data.get("sections"), data.get("opts"),
        tenant_id=t if (t and t != "*") else "")
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True, "id": tid})


@app.put("/api/v1/templates/<tid>")
def templates_update(tid):
    tpl = store.get_template(tid)
    if not tpl:
        return jsonify({"error": "模板不存在"}), 404
    if not _tenant_can_manage_template(tpl):
        return jsonify({"error": "无权修改该模板"}), 403
    data = request.get_json(force=True, silent=True) or {}
    err = store.update_template(
        tid, data.get("name"), data.get("color"), data.get("description"),
        data.get("sections"), data.get("opts"))
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True})


@app.delete("/api/v1/templates/<tid>")
def templates_delete(tid):
    tpl = store.get_template(tid)
    if not tpl:
        return jsonify({"error": "模板不存在"}), 404
    if not _tenant_can_manage_template(tpl):
        return jsonify({"error": "无权删除该模板"}), 403
    err = store.delete_template(tid)
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True})


# ---------------- 文件夹（动态建删，「全部录音」为固定视图不落库） ----------------
@app.get("/api/v1/folders")
def folders_list():
    # P2-1：文件夹全局共享，但计数按租户可见录音统计
    t = _current_tenant()
    return jsonify({"folders": store.list_folders(tenant_id=t if (t and t != "*") else None)})


@app.post("/api/v1/folders")
def folders_create():
    if _current_tenant() not in (None, "*"):  # P2-1：文件夹为全局结构，仅管理员可建
        return jsonify({"error": "仅管理员可创建文件夹"}), 403
    data = request.get_json(force=True, silent=True) or {}
    name, err = store.create_folder(data.get("name"))
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True, "name": name})


@app.delete("/api/v1/folders/<path:name>")
def folders_delete(name):
    if _current_tenant() not in (None, "*"):
        return jsonify({"error": "仅管理员可删除文件夹"}), 403
    moved = store.delete_folder(name)
    return jsonify({"ok": True, "moved_to_inbox": moved})


# ---------------- 录音 ----------------
@app.get("/api/v1/recordings")
def rec_list():
    t = _current_tenant()
    tenant_param = t if (t and t != "*") else None
    # ?deleted=1 → 回收站（已删除，7 天内可恢复）
    deleted = request.args.get("deleted") in ("1", "true")
    items = store.list_recordings(tenant_id=tenant_param, deleted=deleted, limit=500)
    for r in items:
        r["media_url"] = _media_url(r["id"])
    return jsonify({"recordings": items})


@app.get("/api/v1/recordings/<rid>")
def rec_detail(rid):
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    rec = store.get_recording(rid)
    if not rec:
        return jsonify({"error": "录音不存在"}), 404
    job = store.get_job_by_recording(rid)
    result = None
    if job and job.get("result_path"):
        result = _json(job["result_path"])
    # 口语顺滑：去除语气词（仅展示层，不改动落库原文，便于 A/B 对照）
    if result is not None and request.args.get("clean") in ("1", "true"):
        result = polish.clean_result(result)
    rec["media_url"] = _media_url(rid)
    return jsonify({"recording": rec, "job": job, "result": result})


@app.post("/api/v1/recordings/<rid>/turns/edit")
def edit_turn(rid):
    """人工修正某一句转写原文：同步更新结果 JSON（展示/导出）与 segments 表（检索索引）。"""
    if not _tenant_can_access(rid):  # P2-1：修复此前缺失的租户校验（越权编辑面）
        return jsonify({"error": "录音不存在"}), 404
    data = request.get_json(force=True, silent=True) or {}
    try:
        start_ms = int(data.get("start_ms"))
    except (TypeError, ValueError):
        return jsonify({"error": "start_ms 缺失或非法"}), 400
    text = (data.get("text") or "").strip()
    if not text:
        return jsonify({"error": "内容不能为空"}), 400

    job = store.latest_result_job(rid)
    if not job or not job.get("result_path"):
        return jsonify({"error": "该录音暂无转写结果"}), 404
    if job.get("status") not in ("done", "partial"):
        return jsonify({"error": "转写尚未完成，暂不能编辑"}), 400

    result = _json(job["result_path"])
    if not result:
        return jsonify({"error": "结果文件缺失"}), 404

    hit = False
    for t in result.get("turns") or []:
        if t.get("start_ms") == start_ms:
            t["text"] = text
            hit = True
            break
    if not hit:
        return jsonify({"error": "未找到该句（可能已被重新转写，请刷新）"}), 404

    # 同步重建 result["text"]（AI 问答基于该字段），保持与 turns 修改一致
    lm = {x.get("spk"): x.get("label", "说话人?") for x in (result.get("speakers") or [])}
    result["text"] = "\n".join(
        "%s(%s)：%s" % (lm.get(t.get("spk"), "说话人?"), t.get("start", ""), t.get("text", ""))
        for t in result.get("turns") or []
    )
    with open(job["result_path"], "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    n = store.update_segment_text(job["id"], start_ms, t.get("end_ms"), text)
    # 同步刷新检索索引，保证语义搜索命中修正后的文本
    try:
        import search
        search.index_recording_async(rid)  # P0-2：后台重建，编辑接口立即返回
    except Exception:  # noqa: BLE001 索引失败不阻塞编辑
        pass
    # 数据一致性：转写已变，基于旧文本的纪要标记过期、旧润色结果作废
    try:
        store.mark_summaries_stale(rid)
    except Exception:  # noqa: BLE001
        pass
    POLISH_JOBS.pop(rid, None)
    try:
        store.ai_job_del("polish", rid, "")  # 旧润色结果基于已修改的转写，一并作废（P0-1）
    except Exception:  # noqa: BLE001
        pass
    return jsonify({"ok": True, "segments_updated": n})


@app.post("/api/v1/recordings/<rid>/speakers/rename")
def rename_speaker(rid):
    """说话人重命名（P1-1）：改 result JSON 的 speakers[].label，同步全文与检索索引。

    body: {"spk": "1", "label": "张三"}
    副作用范围：本会议内。导出字幕/纪要溯源/搜索结果均随 result 与索引更新。
    已生成的旧纪要文本不回改（与 edit_turn 的 stale 策略一致）。
    """
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    data = request.get_json(force=True, silent=True) or {}
    spk = str(data.get("spk", "")).strip()
    label = (data.get("label") or "").strip()
    if not spk:
        return jsonify({"error": "spk 缺失"}), 400
    if not label:
        return jsonify({"error": "名称不能为空"}), 400
    if len(label) > 30:
        return jsonify({"error": "名称过长（最多 30 字）"}), 400

    job = store.latest_result_job(rid)
    if not job or not job.get("result_path"):
        return jsonify({"error": "该录音暂无转写结果"}), 404
    if job.get("status") not in ("done", "partial"):
        return jsonify({"error": "转写尚未完成，暂不能编辑"}), 400
    result = _json(job["result_path"])
    if not result:
        return jsonify({"error": "结果文件缺失"}), 404

    speakers = result.get("speakers") or []
    hit = False
    for s in speakers:
        if str(s.get("spk")) == spk:
            s["label"] = label
            hit = True
            break
    if not hit:
        return jsonify({"error": "未找到该说话人（可能已被重新转写，请刷新）"}), 404

    # 重写全文（AI 问答基于该字段），与 edit_turn 的组装方式保持一致
    lm = {x.get("spk"): x.get("label", "说话人") for x in speakers}
    result["text"] = "\n".join(
        "%s(%s)：%s" % (lm.get(t.get("spk"), "说话人"), t.get("start", ""), t.get("text", ""))
        for t in result.get("turns") or []
    )
    with open(job["result_path"], "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    # 后台重建检索索引（P0-2）：chunks 携带说话人名，facet 与命中文本随之更新
    try:
        search.index_recording_async(rid)
    except Exception:  # noqa: BLE001 索引失败不阻塞重命名
        pass
    return jsonify({"ok": True, "speakers": speakers})


def _safe_audio_path(p):
    """S2：媒体路径约束——库内路径（DB 被污染/历史脏数据）不得逃出数据目录。

    返回 realpath 后的安全路径；越界返回 None 并记安全日志。
    """
    if not p:
        return None
    real = os.path.realpath(p)
    roots = [os.path.realpath(config.DATA_DIR)]
    for r in roots:
        if real == r or real.startswith(r + os.sep):
            return real
    print("[security] 媒体路径越界已拦截: %s" % p, flush=True)
    return None


@app.get("/api/v1/recordings/<rid>/media")
def rec_media(rid):
    if not _tenant_can_access(rid):
        return jsonify({"error": "音频不存在"}), 404
    if not _media_allowed(rid):
        return jsonify({"error": "invalid or expired media token"}), 401
    norm = os.path.join(config.AUDIO_DIR, rid + ".norm.wav")
    if os.path.isfile(norm):
        return send_file(norm, mimetype="audio/wav", conditional=True)
    rec = store.get_recording(rid)
    safe = _safe_audio_path(rec.get("path")) if rec else None
    if safe and os.path.isfile(safe):
        return send_file(safe, conditional=True)
    return jsonify({"error": "音频不存在"}), 404


# ---------------- 跨会议声纹库（P3-3：仅存本机，可删除，ASR_VOICEPRINT_ENABLED=0 可整体关闭） ----------------
@app.get("/api/v1/voiceprints")
def voiceprints_list():
    t = _current_tenant()
    return jsonify({"voiceprints": store.list_voiceprints(tenant_id=t if (t and t != "*") else None)})


@app.post("/api/v1/voiceprints")
def voiceprints_save():
    """把某录音的某说话人质心保存/合并入声纹库。body {recording_id, spk, name}。"""
    data = request.get_json(silent=True) or {}
    rid = str(data.get("recording_id") or "")
    spk = str(data.get("spk", "")).strip()
    name = (data.get("name") or "").strip()
    if not rid or not spk or not name:
        return jsonify({"error": "recording_id / spk / name 不能为空"}), 400
    if len(name) > 30:
        return jsonify({"error": "名称过长（最多 30 字）"}), 400
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    job = store.latest_result_job(rid)
    if not job:
        return jsonify({"error": "该录音暂无转写结果"}), 400
    try:
        v = store.speaker_centroids_by_global(job["id"]).get(int(spk))
    except ValueError:
        return jsonify({"error": "spk 非法"}), 400
    if v is None:
        return jsonify({"error": "未找到该说话人的声纹特征（可能未开启说话人分离）"}), 404
    t = _current_tenant()
    vid, err = store.save_voiceprint(name, v, tenant_id=t if (t and t != "*") else "", source_rid=rid)
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True, "id": vid})


@app.delete("/api/v1/voiceprints/<vid>")
def voiceprints_delete(vid):
    t = _current_tenant()
    tenant_param = t if (t and t != "*") else None
    if not store.delete_voiceprint(vid, tenant_id=tenant_param):
        return jsonify({"error": "声纹不存在或无权删除"}), 404
    return jsonify({"ok": True})


@app.get("/api/v1/review")
def review_list():
    """复核队列（P2-4）：支持 recording_id / reason / conf_max / offset 筛选。"""
    t = _current_tenant()
    tenant_param = t if (t and t != "*") else None
    try:
        conf_max = request.args.get("conf_max")
        conf_max = float(conf_max) if conf_max not in (None, "") else None
        offset = int(request.args.get("offset", 0) or 0)
        limit = min(int(request.args.get("limit", 200) or 200), 500)
    except ValueError:
        return jsonify({"error": "筛选参数非法"}), 400
    res = store.list_review_segments(
        limit=limit, tenant_id=tenant_param,
        recording_id=request.args.get("recording_id") or None,
        reason=request.args.get("reason") or None,
        conf_max=conf_max, offset=offset)
    return jsonify({"segments": res["items"], "total": res["total"]})


@app.post("/api/v1/review/confirm")
def review_confirm():
    """批量人工复核确认：body {"ids": [segment_id...]}，确认后从复核队列消失。

    P2-1：普通租户仅能确认本租户（含共享）录音的分片。
    """
    data = request.get_json(silent=True) or {}
    ids = data.get("ids") or []
    if not isinstance(ids, list) or not ids:
        return jsonify({"error": "ids 不能为空"}), 400
    t = _current_tenant()
    allowed = None
    if t and t != "*":
        allowed = list(store.list_recording_ids(t))
    confirmed = store.confirm_review_segments(ids, allowed_recording_ids=allowed)
    return jsonify({"ok": True, "confirmed": confirmed})


@app.post("/api/v1/summaries/group")
def summary_group():
    """纪要要点按主题智能分组（LLM）。body {"summary_id", "points":[...]}；按 summary_id 缓存。"""
    data = request.get_json(silent=True) or {}
    sid = str(data.get("summary_id") or "")
    points = data.get("points") or []
    if not sid or not isinstance(points, list) or not points:
        return jsonify({"error": "缺少 summary_id / points"}), 400
    cached = SUMMARY_GROUPS.get(sid)
    if cached is not None:
        return jsonify({"groups": cached, "cached": True})
    try:
        groups = summary.group_points(points)
    except gw.GatewayError as e:
        return jsonify({"error": "模型调用失败：%s" % e}), 502
    except Exception as e:  # noqa: BLE001
        return jsonify({"error": "分组失败：%s" % e}), 500
    SUMMARY_GROUPS[sid] = groups
    return jsonify({"groups": groups, "cached": False})


@app.get("/api/v1/stats")
def stats():
    """可观测性报表：录音/任务/分片状态、队列深度、失败率、平均置信度、复核量。"""
    t = _current_tenant()
    tenant_param = t if (t and t != "*") else None
    return jsonify(store.stats_summary(tenant_id=tenant_param))


# ---------------- 系统信息（用户菜单信息条 / 关于弹层） ----------------
@app.get("/api/v1/system")
def system_info():
    """版本、录音数、音频存储占用、后端运行状态（侧栏系统信息条与关于弹层用）。"""
    n = 0
    try:
        t = _current_tenant()
        tenant_param = t if (t and t != "*") else None
        n = len(store.list_recordings(limit=100000, tenant_id=tenant_param, deleted=False))
    except Exception:
        pass
    size = 0
    try:
        if t and t != "*":
            # P2-1：租户只统计自己（含共享）录音的存储占用
            for r in store.list_recordings(limit=100000, tenant_id=t, deleted=False):
                try:
                    size += os.path.getsize(r["path"]) if r.get("path") else 0
                except OSError:
                    pass
        else:
            for root, _dirs, files in os.walk(config.AUDIO_DIR):
                for fn in files:
                    try:
                        size += os.path.getsize(os.path.join(root, fn))
                    except OSError:
                        pass
    except Exception:
        pass
    return jsonify({
        "version": config.APP_VERSION,
        "brand": "MeetScribe",
        "brand_zh": "聆语会议",
        "recordings": n,
        "storage_bytes": size,
        "status": "running",
        # 安全引导（P0-3）：前端据此展示配置状态与警告
        "auth_configured": bool(config.API_KEY_TENANTS),
        "media_sign_enabled": bool(config.MEDIA_SIGN_SECRET),
        "gateway_configured": bool(getattr(config, "GATEWAY_URL", "")),
    })


# ---------------- AI 网关设置（界面可改，持久化到 runtime_settings.json；P2-1 仅管理员） ----------------
@app.get("/api/v1/settings")
def get_gateway_settings():
    """当前 AI 网关配置（密钥掩码返回）。"""
    denied = _require_admin()
    if denied:
        return denied
    return jsonify(runtime_settings.get())


@app.put("/api/v1/settings")
def put_gateway_settings():
    """更新网关地址 / 密钥 / 模型名，热生效并持久化。"""
    denied = _require_admin()
    if denied:
        return denied
    p = request.get_json(silent=True) or {}
    try:
        cfg, changed = runtime_settings.update(p)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True, "changed": changed, "settings": cfg})


@app.post("/api/v1/settings/test")
def test_gateway_settings():
    """探测网关连通性（GET /models），返回可用模型列表。"""
    denied = _require_admin()
    if denied:
        return denied
    try:
        models = runtime_settings.test()
    except RuntimeError as e:
        return jsonify({"ok": False, "error": str(e)}), 502
    return jsonify({"ok": True, "models": models})


# ---------------- 字幕导出（FR-15，可选） ----------------
def _build_subtitle(result, fmt):
    spk_map = {s["spk"]: s["label"] for s in result.get("speakers", [])}
    def srt_ts(ms):
        ms = max(0, int(ms)); h = ms // 3600000; m = (ms % 3600000) // 60000
        s = (ms % 60000) // 1000; mm = ms % 1000
        return "%02d:%02d:%02d,%03d" % (h, m, s, mm)
    def vtt_ts(ms):
        return srt_ts(ms).replace(",", ".")
    lines = ["WEBVTT\n"] if fmt == "vtt" else []
    for i, t in enumerate(result.get("turns", []), 1):
        label = spk_map.get(t["spk"], "说话人?")
        line = "%s：%s" % (label, t["text"])
        if fmt == "srt":
            lines.append(str(i))
            lines.append("%s --> %s" % (srt_ts(t["start_ms"]), srt_ts(t["end_ms"])))
            lines.append(line)
            lines.append("")
        else:
            lines.append("%s --> %s" % (vtt_ts(t["start_ms"]), vtt_ts(t["end_ms"])))
            lines.append(line)
            lines.append("")
    return "\n".join(lines)


@app.get("/api/v1/recordings/<rid>/subtitle")
def rec_subtitle(rid):
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在或已被删除"}), 404
    rec = store.get_recording(rid)
    if not rec:
        return jsonify({"error": "录音不存在或已被删除"}), 404
    job = store.get_job_by_recording(rid)
    if not job or not job.get("result_path"):
        return jsonify({"error": "该录音还没有转写结果，请等待转写完成后再导出字幕"}), 400
    fmt = (request.args.get("fmt") or request.args.get("format") or "srt").lower()
    if fmt not in ("srt", "vtt"):
        return jsonify({"error": "不支持的字幕格式：%s（可选 srt / vtt）" % fmt}), 400
    result = _json(job["result_path"])
    if not result:
        return jsonify({"error": "转写结果文件缺失或已损坏，请重新转写后再导出"}), 400
    body = _build_subtitle(result, fmt)
    mime = "application/x-subrip" if fmt == "srt" else "text/vtt"
    return Response(body, mimetype=mime,
                    headers={"Content-Disposition": _attachment_header(rec, fmt)})


# ---------------- 导出增强（P1-3）：带时间戳 TXT / Markdown 纪要 ----------------
def _ms_ts(ms):
    ms = max(0, int(ms or 0))
    h, m, s = ms // 3600000, (ms % 3600000) // 60000, (ms % 60000) // 1000
    return "%02d:%02d:%02d" % (h, m, s) if h else "%02d:%02d" % (m, s)


_AUDIO_EXTS = (".wav", ".mp3", ".m4a", ".flac", ".ogg", ".aac", ".wma", ".amr",
               ".mp4", ".mov", ".mkv", ".webm", ".avi")


def _export_base(rec):
    """U4：导出文件名基名 `会议名_YYYYMMDD`，去掉扩展名与非法字符（中文保留）。"""
    name = (rec.get("filename") or "").strip()
    low = name.lower()
    for ext in _AUDIO_EXTS:
        if low.endswith(ext):
            name = name[: -len(ext)]
            break
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name).strip(" ._")
    date = (rec.get("created_at") or "")[:10].replace("-", "")
    base = "_".join(p for p in (name, date) if p)
    return (base or rec.get("id") or "meeting")[:80]


def _attachment_header(rec, ext):
    """Content-Disposition：ASCII 兜底 + RFC5987 中文文件名（浏览器优先用后者）。"""
    name = "%s.%s" % (_export_base(rec), ext)
    fallback = "%s.%s" % (rec.get("id") or "meeting", ext)
    return "attachment; filename=\"%s\"; filename*=UTF-8''%s" % (fallback, quote(name))


def _rec_meta_line(rec, result):
    """导出文件头部的来源信息：录音名 / 日期 / 时长 / 说话人数。"""
    date = (rec.get("created_at") or "")[:10]
    turns = (result or {}).get("turns") or []
    dur = _ms_ts(max((int(x.get("end_ms") or 0) for x in turns), default=0))
    parts = []
    if rec.get("filename"):
        parts.append("录音文件：%s" % rec["filename"])
    if date:
        parts.append("日期：%s" % date)
    parts.append("时长：%s" % dur)
    if (result or {}).get("speakers"):
        parts.append("说话人：%d 位" % len(result["speakers"]))
    return "> " + " · ".join(parts)


def _build_txt(result):
    """带时间戳的纯文本转写：`说话人 [mm:ss] 内容`（P1-3）。"""
    spk_map = {s["spk"]: s["label"] for s in result.get("speakers", [])}
    return "\n".join(
        "%s [%s] %s" % (spk_map.get(t["spk"], "说话人?"), _ms_ts(t.get("start_ms")), t.get("text", ""))
        for t in result.get("turns", [])
    )


def _md_points(points):
    out = []
    for p in points or []:
        src = "（%s%s）" % (p.get("speaker", ""), (" " + p["timestamp"]) if p.get("timestamp") else "") \
            if (p.get("speaker") or p.get("timestamp")) else ""
        out.append("- %s%s" % (p.get("content", ""), src))
    return out


def _summary_to_markdown(content, template, title, meta=None):
    """把纪要 JSON 转成 Markdown（标题/一句话/要点/行动项/风险/待办，均带时间戳溯源）。"""
    c = content or {}
    lines = ["# %s" % (c.get("title") or title)]
    if meta:
        lines.append("\n%s" % meta)
    if c.get("one_line"):
        lines.append("\n> %s" % c["one_line"])
    if c.get("owner"):
        lines.append("\n**备忘人**：%s" % c["owner"])
    if c.get("attendees"):
        lines.append("\n**参与人**：%s" % "、".join(c["attendees"]))
    sections = c.get("sections")
    if isinstance(sections, dict):  # 自定义模板：sections 分组
        for sec, pts in sections.items():
            lines.append("\n## %s" % sec)
            lines.extend(_md_points(pts))
    elif template == "memo":
        lines.append("\n## 要点")
        lines.extend(_md_points(c.get("points")))
        if c.get("todos"):
            lines.append("\n## 待办")
            lines.extend("-[ ] %s%s%s" % (a.get("task", ""),
                                         " @%s" % a["owner"] if a.get("owner") else "",
                                         "（截止 %s）" % a["deadline"] if a.get("deadline") else "")
                         for a in c["todos"])
    else:
        lines.append("\n## 讨论要点")
        lines.extend(_md_points(c.get("points")))
        if c.get("action_items"):
            lines.append("\n## 行动项")
            lines.extend("-[ ] %s%s%s" % (a.get("task", ""),
                                         " @%s" % a["owner"] if a.get("owner") else "",
                                         "（截止 %s）" % a["deadline"] if a.get("deadline") else "")
                         for a in c["action_items"])
        if c.get("risks"):
            lines.append("\n## 风险")
            lines.extend("- %s%s" % (r.get("content", ""),
                                     " [%s]" % r["level"] if r.get("level") else "")
                         for r in c["risks"])
    return "\n".join(lines) + "\n"


def _transcript_markdown(rec, result):
    lines = ["# %s" % rec.get("filename", rec.get("id", "")), "",
             _rec_meta_line(rec, result), "",
             "> 转写全文（带时间戳，尚未生成纪要）", ""]
    lines.extend(_build_txt(result).split("\n"))
    return "\n".join(lines) + "\n"


# ---------------- B4 Word(.docx) 导出 ----------------
def _docx_todos(doc, items):
    for a in items or []:
        line = "□ %s" % a.get("task", "")
        if a.get("owner"):
            line += " @%s" % a["owner"]
        if a.get("deadline"):
            line += "（截止 %s）" % a["deadline"]
        doc.add_paragraph(line, style="List Bullet")


def _build_docx(s, template, rec, result):
    """纪要转 Word（有纪要时含结构化章节），并附全文转写（带时间戳）。"""
    import io

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    title = rec.get("filename") or rec.get("id") or "会议纪要"
    if s and s.get("content"):
        c = s["content"]
        doc.add_heading(c.get("title") or title, level=0)
        if c.get("one_line"):
            p = doc.add_paragraph()
            r = p.add_run(c["one_line"])
            r.bold = True
        if c.get("owner"):
            doc.add_paragraph("备忘人：%s" % c["owner"])
        if c.get("attendees"):
            doc.add_paragraph("参与人：%s" % "、".join(c["attendees"]))

        def add_points(pts):
            for p in pts or []:
                src = ""
                if p.get("speaker") or p.get("timestamp"):
                    src = "（%s %s）" % (p.get("speaker", ""), p.get("timestamp", ""))
                doc.add_paragraph("%s%s" % (p.get("content", ""), src), style="List Bullet")

        sections = c.get("sections")
        if isinstance(sections, dict):  # 自定义模板
            for sec, pts in sections.items():
                doc.add_heading(sec, level=1)
                add_points(pts)
        elif template == "memo":
            doc.add_heading("要点", level=1)
            add_points(c.get("points"))
            if c.get("todos"):
                doc.add_heading("待办", level=1)
                _docx_todos(doc, c["todos"])
        else:
            doc.add_heading("讨论要点", level=1)
            add_points(c.get("points"))
            if c.get("action_items"):
                doc.add_heading("行动项", level=1)
                _docx_todos(doc, c["action_items"])
            if c.get("risks"):
                doc.add_heading("风险", level=1)
                for r in c["risks"]:
                    lv = " [%s]" % r["level"] if r.get("level") else ""
                    doc.add_paragraph("%s%s" % (r.get("content", ""), lv), style="List Bullet")
    else:
        doc.add_heading(title, level=0)
        doc.add_paragraph("转写全文（带时间戳，尚未生成纪要）")

    doc.add_heading("转写全文", level=1)
    spk_map = {x["spk"]: x["label"] for x in result.get("speakers", [])}
    for t in result.get("turns", []):
        p = doc.add_paragraph()
        r = p.add_run("[%s] %s：" % (_ms_ts(t.get("start_ms")), spk_map.get(t["spk"], "说话人?")))
        r.font.size = Pt(9)
        p.add_run(t.get("text", ""))

    # U4：写入文档属性，接收方在 Word「文件 → 信息」里能看到来源
    cp = doc.core_properties
    cp.title = (s["content"].get("title") if s and s.get("content") else None) or title
    cp.author = "MeetScribe"
    cp.last_modified_by = "MeetScribe"
    cp.subject = "会议纪要" if (s and s.get("content")) else "转写全文"
    cp.comments = "由 MeetScribe 导出 · 录音 %s" % (rec.get("id") or "")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@app.get("/api/v1/recordings/<rid>/export")
def rec_export(rid):
    """P1-3 导出：fmt=txt 带时间戳纯文本转写；fmt=md Markdown 纪要（无纪要时回退转写全文）。

    U3：所有失败路径都返回可直接展示给用户的中文原因，
    前端（utils/download）会读 error 字段并弹提示，避免浏览器把 JSON 当成文件下载。
    """
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在或已被删除"}), 404
    rec = store.get_recording(rid)
    if not rec:
        return jsonify({"error": "录音不存在或已被删除"}), 404
    job = store.get_job_by_recording(rid)
    if not job or not job.get("result_path"):
        return jsonify({"error": "该录音还没有转写结果，请等待转写完成后再导出"}), 400
    result = _json(job["result_path"])
    if not result:
        return jsonify({"error": "转写结果文件缺失或已损坏，请重新转写后再导出"}), 400
    fmt = (request.args.get("fmt") or request.args.get("format") or "txt").lower()
    if fmt not in ("txt", "md", "docx"):
        return jsonify({"error": "不支持的导出格式：%s（可选 txt / md / docx）" % fmt}), 400
    if fmt == "docx":
        s = store.get_summary(rid)
        if s and s.get("status") != "done":
            s = None
        try:
            body = _build_docx(s, (s.get("template") if s else None) or "meeting", rec, result)
        except ImportError:
            return jsonify({"error": "Word 导出不可用：服务端缺少 python-docx 依赖，"
                                     "请改用 Markdown 或 TXT 导出"}), 500
        except Exception as e:  # noqa: BLE001
            print("[export] docx 生成失败 rid=%s: %r" % (rid, e), flush=True)
            return jsonify({"error": "Word 导出失败，请改用 Markdown 或 TXT 导出"}), 500
        return Response(body,
                        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": _attachment_header(rec, "docx")})
    if fmt == "md":
        s = store.get_summary(rid)
        if s and s.get("content") and s.get("status") == "done":
            body = _summary_to_markdown(s["content"], s.get("template") or "meeting",
                                        rec.get("filename") or rid, _rec_meta_line(rec, result))
        else:
            body = _transcript_markdown(rec, result)
        mime, ext = "text/markdown", "md"
    else:
        fmt = "txt"
        body = _build_txt(result)
        mime, ext = "text/plain", "txt"
    return Response(body, mimetype=mime,
                    headers={"Content-Disposition": _attachment_header(rec, ext)})


# ---------------- 历史语义检索（FR-25~FR-32） ----------------
@app.post("/api/v1/search")
def search_api():
    data = request.get_json(silent=True) or {}
    q = (data.get("query") or request.form.get("query") or "").strip()
    filters = {
        "date_from": data.get("date_from"),
        "date_to": data.get("date_to"),
        "speaker": data.get("speaker"),
        "recording_id": data.get("recording_id"),
    }
    # FR-31 检索层租户隔离：非管理员密钥仅能在自身租户（含共享）录音内检索
    t = _current_tenant()
    tenant_ids = store.list_recording_ids(t) if (t and t != "*") else None
    # P1-4 文件夹过滤：'' = 未归档；与租户范围求交集，双重防护越权
    folder = data.get("folder")
    if folder is not None:
        ids = store.list_recording_ids(t, folder=str(folder))
        if tenant_ids is not None:
            ids &= tenant_ids
        filters["recording_ids"] = list(ids)
    elif tenant_ids is not None:
        filters["recording_ids"] = list(tenant_ids)
    page = int(data.get("page", 1) or 1)
    page_size = min(int(data.get("page_size", 10) or 10), 100)
    return jsonify(search.search(q, filters, page, page_size))


@app.post("/api/v1/recordings/<rid>/index")
def rec_index(rid):
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    if not store.get_recording(rid):
        return jsonify({"error": "录音不存在"}), 404
    # P0-2：改为后台异步重建，立即返回（状态可通过 /search 查询 unindexed 观察）
    search.index_recording_async(rid)
    return jsonify({"ok": True, "started": True})


# P0-2 全量重建：后台执行 + 进度查询（原同步实现会长时间挂住 HTTP 请求）
_REINDEX = {"running": False, "done": 0, "failed": 0, "total": 0,
            "started_at": None, "finished_at": None, "error": None}


@app.post("/api/v1/search/reindex")
def search_reindex():
    denied = _require_admin()  # P2-1：全量重建影响全局，仅管理员
    if denied:
        return denied
    if _REINDEX["running"]:
        return jsonify({"ok": True, "started": False, "status": _REINDEX})
    _REINDEX.update(running=True, done=0, failed=0, total=0,
                    started_at=time.strftime("%Y-%m-%d %H:%M:%S"),
                    finished_at=None, error=None)

    def _run():
        try:
            ok, skip = search.index_all_done(
                progress=lambda done, failed, total: _REINDEX.update(
                    done=done, failed=failed, total=total))
            _REINDEX.update(running=False, finished_at=time.strftime("%Y-%m-%d %H:%M:%S"))
            print("[search] 全量重建完成：成功=%d 跳过=%d" % (ok, skip), flush=True)
        except Exception as e:  # noqa: BLE001
            _REINDEX.update(running=False, error=str(e)[:500],
                            finished_at=time.strftime("%Y-%m-%d %H:%M:%S"))

    threading.Thread(target=_run, daemon=True).start()
    return jsonify({"ok": True, "started": True, "status": _REINDEX})


@app.get("/api/v1/search/reindex/status")
def search_reindex_status():
    return jsonify({"status": _REINDEX})


@app.get("/api/v1/search/logs")
def search_logs():
    denied = _require_admin()  # P2-1：检索审计含全部租户查询词，仅管理员
    if denied:
        return denied
    return jsonify({"logs": search.recent_search_logs()})


# ---------------- 可观测性（P2-2）：任务事件流水与运行指标 ----------------
@app.get("/api/v1/recordings/<rid>/events")
def rec_events(rid):
    """任务事件流水（分片耗时/重试/失败原因），租户隔离。"""
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    job = store.get_job_by_recording(rid)
    if not job:
        return jsonify({"events": []})
    return jsonify({"events": store.job_events(job["id"])})


@app.get("/api/v1/system/metrics")
def system_metrics():
    """运维指标：队列、进行中任务、24h 成功率、最近失败事件（P2-2，仅管理员）。"""
    denied = _require_admin()
    if denied:
        return denied
    s = store.stats_summary(tenant_id=None)
    return jsonify({
        "queue": s.get("queue"),
        "jobs": s.get("jobs"),
        "jobs_today": s.get("jobs_today"),
        "chunks": s.get("chunks"),
        "segments": {k: s.get("segments", {}).get(k) for k in ("count", "avg_confidence", "need_review")},
        "recent_failures": store.recent_failed_events(limit=20),
        "reindex": _REINDEX,
        "uptime_s": int(time.time() - _START_TIME),
    })


@app.get("/api/v1/system/metrics.txt")
def metrics_prometheus():
    """Prometheus 文本格式指标（C2，仅管理员）。

    抓取配置示例（prometheus.yml）：
      scrape_configs:
        - job_name: meetscribe
          metrics_path: /api/v1/system/metrics.txt
          static_configs:
            - targets: ["host:5000"]
    （Bearer 鉴权场景在 scrape_config 加 authorization.credentials）
    """
    denied = _require_admin()
    if denied:
        return denied
    s = store.stats_summary(tenant_id=None)
    L = []

    def g(name, mtype, help_text, val, labels=""):
        L.append("# HELP %s %s" % (name, help_text))
        L.append("# TYPE %s %s" % (name, mtype))
        L.append("%s%s %s" % (name, labels, val))

    g("meetscribe_uptime_seconds", "gauge", "Process uptime in seconds",
      int(time.time() - _START_TIME))
    for st, n in (s.get("jobs") or {}).items():
        g("meetscribe_jobs", "gauge", "Jobs by status", n, '{status="%s"}' % st)
    for st, n in (s.get("recordings") or {}).items():
        g("meetscribe_recordings", "gauge", "Recordings by status", n, '{status="%s"}' % st)
    q = s.get("queue") or {}
    for k in ("queued", "processing"):
        g("meetscribe_queue", "gauge", "Job queue depth by state", q.get(k, 0), '{state="%s"}' % k)
    jt = s.get("jobs_today") or {}
    g("meetscribe_jobs_today_total", "counter", "Jobs created today", jt.get("count", 0))
    g("meetscribe_jobs_today_failed", "counter", "Jobs failed today", jt.get("failed", 0))
    seg = s.get("segments") or {}
    if seg.get("avg_confidence") is not None:
        g("meetscribe_segments_avg_confidence", "gauge", "Average segment confidence",
          seg["avg_confidence"])
    g("meetscribe_segments_need_review", "gauge", "Segments pending review",
      seg.get("need_review", 0))
    for st, n in (s.get("chunks") or {}).items():
        g("meetscribe_chunks", "gauge", "Chunks by status", n, '{status="%s"}' % st)
    return Response("\n".join(L) + "\n", mimetype="text/plain; version=0.0.4; charset=utf-8")


# ---------------- AI 问答（FR-19 延伸） ----------------
@app.post("/api/v1/ask")
def ask_api():
    data = request.get_json(silent=True) or {}
    q = (data.get("question") or "").strip()
    if not q:
        return jsonify({"error": "缺少 question"}), 400
    rid = data.get("recording_id")
    rids = list(data.get("recording_ids") or [])
    if rid:
        rids = [rid] + [r for r in rids if r != rid]
    if not rids:
        # 跨录音模式：默认取最近已完成录音（最多 10 条）
        rids = [r["id"] for r in store.list_recordings(limit=50) if r.get("status") == "done"][:10]
    t = _current_tenant()
    if t and t != "*":
        allowed = set(store.list_recording_ids(t))
        rids = [r for r in rids if r in allowed][:10]
    else:
        rids = rids[:10]
    history = data.get("history") or []
    try:
        ans, used = ask.answer(rids, q, history)
    except gw.GatewayError as e:
        return jsonify({"error": "模型调用失败：%s" % e}), 502
    return jsonify({"answer": ans, "recordings": used, "model": config.GATEWAY_LLM_MODEL})


def _purge_recording(rec):
    """删除录音及其全部关联数据：DB 级联 + 向量索引 + 音频/结果文件（FR-30/34）。"""
    rid = rec["id"]
    jid = store.delete_recording(rid)
    try:
        search.delete_index(rid)  # 同步清理检索索引
    except Exception:  # noqa: BLE001
        pass
    for p in (rec.get("path"), os.path.join(config.AUDIO_DIR, rid + ".norm.wav")):
        try:
            if p and os.path.isfile(p):
                os.remove(p)
        except OSError:
            pass
    if jid:
        rp = os.path.join(config.RESULT_DIR, jid + ".json")
        try:
            if os.path.isfile(rp):
                os.remove(rp)
        except OSError:
            pass
    return jid


@app.delete("/api/v1/recordings/<rid>")
def rec_delete(rid):
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    rec = store.get_recording(rid)
    if not rec:
        return jsonify({"error": "录音不存在"}), 404
    if request.args.get("forever") in ("1", "true"):
        _purge_recording(rec)               # 彻底删除（回收站清除，不可恢复）
    else:
        store.soft_delete_recording(rid)    # 默认移入回收站，7 天可恢复
        try:
            import search
            search.delete_index(rid)        # 同步移除语义索引，回收站内容不再可被检索
        except Exception:  # noqa: BLE001
            pass
    return jsonify({"ok": True})


@app.post("/api/v1/recordings/<rid>/restore")
def rec_restore(rid):
    """从回收站恢复录音。"""
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    store.restore_recording(rid)
    try:
        import search
        search.index_recording_async(rid)   # 恢复后后台重建语义索引（P0-2）
    except Exception:  # noqa: BLE001
        pass
    return jsonify({"ok": True})


@app.patch("/api/v1/recordings/<rid>")
def rec_update(rid):
    """更新录音元信息：目前支持 folder（移动到文件夹，'' 表示不归类）。"""
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    data = request.get_json(silent=True) or {}
    if "folder" in data:
        store.set_recording_folder(rid, str(data.get("folder") or ""))
    return jsonify({"ok": True})


@app.post("/api/v1/recordings/batch-folder")
def rec_batch_folder():
    """批量移动到文件夹：body {"ids": [...], "folder": "项目复盘"}，单次上限 100。"""
    data = request.get_json(silent=True) or {}
    ids = data.get("ids") or []
    folder = str(data.get("folder") or "")
    if not isinstance(ids, list) or not ids:
        return jsonify({"error": "ids 不能为空"}), 400
    moved = 0
    for rid in ids[:100]:
        if _tenant_can_access(rid) and store.get_recording(rid):
            store.set_recording_folder(rid, folder)
            moved += 1
    return jsonify({"ok": True, "moved": moved})


@app.post("/api/v1/recordings/batch-delete")
def rec_batch_delete():
    """FR-33 批量删除：body {"ids": [...], "forever": bool(默认 false=移入回收站)}，单次上限 100。"""
    data = request.get_json(silent=True) or {}
    ids = data.get("ids") or []
    forever = bool(data.get("forever"))
    if not isinstance(ids, list) or not ids:
        return jsonify({"error": "ids 不能为空"}), 400
    results = []
    for rid in ids[:100]:
        if not _tenant_can_access(rid):
            results.append({"id": rid, "ok": False, "error": "录音不存在或无权访问"})
            continue
        rec = store.get_recording(rid)
        if not rec:
            results.append({"id": rid, "ok": False, "error": "录音不存在"})
            continue
        if forever:
            _purge_recording(rec)
        else:
            store.soft_delete_recording(rid)
        results.append({"id": rid, "ok": True})
    return jsonify({"results": results,
                    "deleted": sum(1 for r in results if r["ok"])})


# ---------------- 智能润色（LLM 合并碎句 / 去语气词 / 书面化） ----------------
POLISH_JOBS = {}  # rid -> {"status": processing/done/failed, "result": ..., "error": ...}
SUMMARY_GROUPS = {}  # summary_id -> [{"title": 主题, "idx": [要点下标...]}]（纪要要点主题分组缓存）


def _polish_worker(rid):
    key = rid
    try:
        res = rewrite.polish_transcript(rid)
        POLISH_JOBS[key] = {"status": "done", "result": res}
        # 写库持久化（P0-1）：重启后结果不丢
        try:
            store.ai_job_upsert("polish", rid, "", status="done",
                                result_json=json.dumps(res, ensure_ascii=False))
        except Exception:  # noqa: BLE001
            pass
    except Exception as e:  # noqa: BLE001
        POLISH_JOBS[key] = {"status": "failed", "error": str(e)}
        try:
            store.ai_job_upsert("polish", rid, "", status="failed", error=str(e))
        except Exception:  # noqa: BLE001
            pass


@app.post("/api/v1/recordings/<rid>/polish")
def gen_polish(rid):
    """提交 LLM 润色任务，立即返回（约 1~2 分钟，前端轮询 GET）。"""
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    if not store.get_recording(rid):
        return jsonify({"error": "录音不存在"}), 404
    if not store.get_job_by_recording(rid):
        return jsonify({"error": "该录音无转写任务，无法润色"}), 400
    # 原子占位：并发双击只起一个 LLM 润色线程
    with _IDEM_LOCK:
        st = POLISH_JOBS.get(rid)
        if st and st["status"] == "processing":
            return jsonify({"status": "processing"})
        POLISH_JOBS[rid] = {"status": "processing"}
    try:
        store.ai_job_upsert("polish", rid, "", status="processing")
    except Exception:  # noqa: BLE001
        pass
    threading.Thread(target=_polish_worker, args=(rid,), daemon=True).start()
    return jsonify({"status": "processing"})


@app.get("/api/v1/recordings/<rid>/polish")
def get_polish(rid):
    if not _tenant_can_access(rid):
        return jsonify({"status": "not_started"})
    st = POLISH_JOBS.get(rid)
    if st and st["status"] == "processing":
        return jsonify({"status": "processing"})
    if st and st["status"] == "failed":
        return jsonify({"status": "failed", "error": st.get("error")})
    if st and st["status"] == "done":
        return jsonify({"status": "done", "result": st["result"]})
    # 内存无记录（服务重启过）→ 回落数据库（P0-1）
    db = store.ai_job_get("polish", rid, "")
    if not db:
        return jsonify({"status": "not_started"})
    if db["status"] == "processing":
        return jsonify({"status": "processing"})
    if db["status"] == "failed":
        return jsonify({"status": "failed", "error": db.get("error")})
    if db["status"] == "done":
        try:
            return jsonify({"status": "done", "result": json.loads(db["result_json"] or "null")})
        except Exception:  # noqa: BLE001
            return jsonify({"status": "not_started"})
    return jsonify({"status": "not_started"})


# ---------------- 文本总结（FR-19~FR-24，异步） ----------------
SUMMARY_JOBS = {}  # (rid, template) -> {"status": processing/done/failed, ...}
_IDEM_LOCK = threading.Lock()  # complete 幂等检查与写入的原子性（防并发双建录音/任务）


def _summary_worker(rid, template):
    key = (rid, template)
    try:
        res = summary.summarize(rid, template)
        SUMMARY_JOBS[key] = {"status": "done", "summary_id": res["_id"],
                             "model": res["_model"], "version": res["_version"]}
        # 写库持久化（P0-1）
        try:
            store.ai_job_upsert("summary", rid, template, status="done",
                                summary_id=res["_id"], model=res["_model"],
                                version=res["_version"])
        except Exception:  # noqa: BLE001
            pass
    except Exception as e:  # noqa: BLE001
        SUMMARY_JOBS[key] = {"status": "failed", "error": str(e)}
        try:
            store.ai_job_upsert("summary", rid, template, status="failed", error=str(e))
        except Exception:  # noqa: BLE001
            pass


@app.post("/api/v1/recordings/<rid>/summary")
def gen_summary(rid):
    """提交总结任务，立即返回（LLM 生成约 1~3 分钟，前端轮询 GET）。"""
    data = request.get_json(silent=True) or {}
    template = data.get("template") or request.form.get("template") or "meeting"
    if template not in ("meeting", "memo") and not store.get_template(template):
        return jsonify({"error": "模板不存在或已被删除，请刷新模板列表"}), 400
    if not _tenant_can_access(rid):
        return jsonify({"error": "录音不存在"}), 404
    if not store.get_recording(rid):
        return jsonify({"error": "录音不存在"}), 404
    if not store.get_job_by_recording(rid):
        return jsonify({"error": "该录音无转写任务，无法总结"}), 400
    force = str(data.get("force", "")).lower() in ("1", "true")
    key = (rid, template)
    # 原子占位：并发重复提交只起一个 LLM 线程
    with _IDEM_LOCK:
        st = SUMMARY_JOBS.get(key)
        if st and st["status"] == "processing":
            return jsonify({"status": "processing"})
        if store.get_summary(rid, template) and not force:
            return jsonify({"status": "done", "cached": True})
        SUMMARY_JOBS[key] = {"status": "processing"}
    try:
        store.ai_job_upsert("summary", rid, template, status="processing")
    except Exception:  # noqa: BLE001
        pass
    threading.Thread(target=_summary_worker, args=(rid, template), daemon=True).start()
    return jsonify({"status": "processing"})


@app.get("/api/v1/recordings/<rid>/summary")
def get_summary(rid):
    if not _tenant_can_access(rid):
        return jsonify({"status": "not_started"})
    template = request.args.get("template") or "meeting"
    key = (rid, template)
    st = SUMMARY_JOBS.get(key)
    if st and st["status"] == "processing":
        return jsonify({"status": "processing"})
    # force 重生成失败时优先暴露 failed，而不是把 DB 里的旧纪要伪装成 done
    if st and st["status"] == "failed":
        return jsonify({"status": "failed", "error": st.get("error")})
    # 内存无记录（服务重启过）→ 回落 ai_jobs 表（P0-1）
    if not st:
        db = store.ai_job_get("summary", rid, template)
        if db and db["status"] == "processing":
            return jsonify({"status": "processing"})
        if db and db["status"] == "failed":
            return jsonify({"status": "failed", "error": db.get("error")})
    s = store.get_summary(rid, template)
    if not s:
        return jsonify({"status": "not_started"})
    out = dict(s)
    out["status"] = "done"
    return jsonify(out)


@app.post("/api/v1/summaries/<sid>/edit")
def edit_summary(sid):
    s = store.get_summary_by_id(sid)
    if not s:
        return jsonify({"error": "摘要不存在"}), 404
    if not _tenant_can_access(s["recording_id"]):
        return jsonify({"error": "摘要不存在"}), 404
    data = request.get_json(silent=True) or {}
    editor = data.get("editor") or request.form.get("editor") or "anonymous"
    before = data.get("before") or ""
    after = data.get("after") or ""
    store.save_summary_edit(sid, editor, before, after)
    return jsonify({"ok": True})


# ---------------- 同步转写（短音频自测） ----------------
@app.post("/api/v1/transcribe")
def transcribe_sync():
    f = request.files.get("audio")
    if f is None or f.filename == "":
        return jsonify({"error": "未收到音频文件"}), 400
    ext = os.path.splitext(f.filename)[1] or ".wav"
    fd, tmp = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    f.save(tmp)
    hotword = request.form.get("hotword", "")
    preset = request.form.get("preset_spk_num") or request.form.get("num_speakers")
    preset = int(preset) if preset and int(preset) > 0 else None
    try:
        wav = os.path.join(config.AUDIO_DIR, "sync_" + os.path.basename(tmp) + ".wav")
        eng = _engine()
        eng.decode_to_wav(tmp, wav)
        dur = eng.probe_duration_ms(wav)
        res = eng.process_chunk(wav, 0, dur, wav + ".c.wav", hotword, preset)
        for s in res["sentences"]:
            s["start"] = eng.fmt_ms(s["start_ms"])
            s["end"] = eng.fmt_ms(s["end_ms"])
            s["spk"] = s["local_spk"]
        text = "".join(s["text"] for s in res["sentences"])
        return jsonify({"text": text, "sentences": res["sentences"],
                        "segments": len(res["sentences"]), "duration_ms": dur})
    finally:
        for p in (tmp,):
            try:
                os.remove(p)
            except OSError:
                pass


def _purge_expired_and_orphans():
    """回收站 7 天自动清理（FR-34）+ 上传临时目录 24h 孤儿清理。"""
    try:
        for rec in store.list_expired_deleted(7):
            print("[purge] 回收站过期，自动清理:", rec["id"], rec.get("filename"), flush=True)
            _purge_recording(rec)
    except Exception as e:  # noqa: BLE001 清理失败不阻塞服务
        print("[purge] 回收站清理异常:", e, flush=True)
    try:
        now = time.time()
        for name in os.listdir(config.UPLOAD_DIR):
            d = os.path.join(config.UPLOAD_DIR, name)
            if os.path.isdir(d) and now - os.path.getmtime(d) > 24 * 3600:
                import shutil
                shutil.rmtree(d, ignore_errors=True)
                print("[purge] 清理孤儿上传目录:", name, flush=True)
    except Exception as e:  # noqa: BLE001
        pass


def _start_housekeeping():
    """后台守护线程：启动时清理一次，此后每 6 小时一次。"""
    def _loop():
        _purge_expired_and_orphans()
        while True:
            time.sleep(6 * 3600)
            _purge_expired_and_orphans()
    threading.Thread(target=_loop, daemon=True, name="housekeeping").start()


if __name__ == "__main__":
    worker.start_workers()
    _start_housekeeping()
    # S3：公网"裸奔"检测——监听非本机且未配置 API 鉴权时打醒目警告
    _host = os.environ.get("ASR_HOST", "0.0.0.0")
    if _host not in ("127.0.0.1", "localhost", "::1") and not config.API_KEY_TENANTS:
        print("=" * 64, flush=True)
        print("⚠️  安全警告：服务监听 %s，但未配置 ASR_API_KEYS——" % _host, flush=True)
        print("    同一网络内任何人可上传录音并读取全部数据。", flush=True)
        print("    公网/办公网部署请配置 ASR_API_KEYS 与 ASR_MEDIA_SIGN_SECRET，", flush=True)
        print("    示例见仓库 .env.example。", flush=True)
        print("=" * 64, flush=True)
    print(">> ASR 服务启动：http://0.0.0.0:5000  (API 前缀 /api/v1)", flush=True)
    app.run(host="0.0.0.0", port=5000, threaded=True, debug=False, use_reloader=False)
